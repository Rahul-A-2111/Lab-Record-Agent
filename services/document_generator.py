"""
Document Generation Service.

Builds publication-grade collegiate Word (.docx) lab record documents using python-docx.
Supports customized borders, typography (separate heading & content sizes),
monospaced source code listings, embedded screenshots/figures, and evaluator signature blocks.
"""

import base64
import io
import os
import re
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def generate_docx(
    output_path: str,
    sections: dict,
    formatting: dict | None = None,
    images: list | None = None,
    meta: dict | None = None,
) -> str:
    """
    Creates a formal collegiate lab record Word document (.docx).

    Args:
        output_path: Destination file path for saving the document.
        sections: Dictionary containing AIM, ALGORITHM, PROGRAM, OUTPUT, RESULT content.
        formatting: Border styles, typography settings, line spacing, alignments.
        images: List of figure dictionaries containing image paths or base64 data.
        meta: Title, experiment number, date, department details.

    Returns:
        output_path (str): The verified path of the generated .docx file.
    """
    formatting = formatting or {}
    images = images or []
    meta = meta or {}

    font_family = formatting.get("font_family", "Times New Roman")
    heading_size = float(formatting.get("heading_size", 14))
    content_size = float(formatting.get("content_size", 11))
    line_spacing = float(formatting.get("line_spacing", 1.15))
    border_style = formatting.get("border_style", "single").lower()
    border_thickness = float(formatting.get("border_thickness", 1.0))
    border_color = formatting.get("border_color", "#000000").lstrip("#")
    body_alignment = formatting.get("alignment", "justify").lower()

    align_map = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    body_align_val = align_map.get(body_alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc = docx.Document()

    # 1. Setup A4 Page Layout & Margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)  # Left gutter for record binding
    section.right_margin = Inches(1.0)

    # 2. Setup Page Borders if requested
    if border_style in ("single", "double"):
        sz = str(int(border_thickness * 8))  # 1/8 pt units
        border_xml = (
            f'<w:pgBorders {nsdecls("w")}>'
            f'<w:top w:val="{border_style}" w:sz="{sz}" w:space="24" w:color="{border_color}"/>'
            f'<w:left w:val="{border_style}" w:sz="{sz}" w:space="24" w:color="{border_color}"/>'
            f'<w:bottom w:val="{border_style}" w:sz="{sz}" w:space="24" w:color="{border_color}"/>'
            f'<w:right w:val="{border_style}" w:sz="{sz}" w:space="24" w:color="{border_color}"/>'
            f'</w:pgBorders>'
        )
        section._sectPr.append(parse_xml(border_xml))

    # 3. Institutional Header
    dept_name = meta.get("department", "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    lab_name = meta.get("lab_name", "Laboratory Record Manual")

    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_after = Pt(2)
    r_dept = p_header.add_run(dept_name.upper())
    r_dept.font.name = font_family
    r_dept.font.size = Pt(heading_size + 2)
    r_dept.bold = True

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run(lab_name)
    r_sub.font.name = font_family
    r_sub.font.size = Pt(content_size)
    r_sub.italic = True

    # 4. Experiment Meta Strip
    exp_no = meta.get("exp_no", "04")
    exp_date = meta.get("date", "OCTOBER 24, 2024")
    title = meta.get("title", "IMPLEMENTATION OF PROGRAM")

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(8)
    r_exp = p_meta.add_run(f"EXPERIMENT NO: {exp_no}")
    r_exp.font.name = font_family
    r_exp.font.size = Pt(content_size)
    r_exp.bold = True

    # Right align date via tabs
    p_meta.add_run(f"\t\t\t\tDATE: {exp_date}")

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(16)
    r_title = p_title.add_run(title.upper())
    r_title.font.name = font_family
    r_title.font.size = Pt(heading_size + 1)
    r_title.bold = True
    r_title.underline = True

    # Helper function for adding headings
    def add_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(f"{text}:")
        run.font.name = font_family
        run.font.size = Pt(heading_size)
        run.bold = True
        run.underline = True
        return p

    def get_align_val(align_str: str | None):
        if not align_str:
            return body_align_val
        a = align_str.lower()
        if a == "left":
            return WD_ALIGN_PARAGRAPH.LEFT
        elif a == "center":
            return WD_ALIGN_PARAGRAPH.CENTER
        elif a == "justify":
            return WD_ALIGN_PARAGRAPH.JUSTIFY
        return body_align_val

    # Normalize incoming sections into an ordered list: [(key, title, content, alignment)]
    ordered_items = []
    if isinstance(sections, list):
        for item in sections:
            if isinstance(item, dict):
                k = item.get("key", "").upper()
                t = item.get("title", k)
                c = item.get("content", "")
                align = item.get("alignment")
                ordered_items.append((k, t, c, align))
    elif isinstance(sections, dict):
        # If section_order is passed in payload or meta, respect it
        s_order = meta.get("section_order") if meta else None
        if not s_order:
            s_order = list(sections.keys())
        for k in s_order:
            if k in sections and sections[k]:
                c = sections[k]
                align = None
                if isinstance(c, dict):
                    content_val = c.get("content", "")
                    align = c.get("alignment")
                else:
                    content_val = c
                ordered_items.append((k.upper(), k.upper(), content_val, align))

    # 5. Output Selected Sections in the EXACT requested order without boxes
    for key, title_text, content_val, item_align in ordered_items:
        sec_align = get_align_val(item_align)

        if key == "AIM":
            add_heading(title_text)
            p = doc.add_paragraph()
            p.alignment = sec_align
            p.paragraph_format.line_spacing = line_spacing
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(str(content_val).strip())
            run.font.name = font_family
            run.font.size = Pt(content_size)

        elif key == "ALGORITHM":
            add_heading(title_text)
            if isinstance(content_val, list):
                steps = content_val
            else:
                steps = str(content_val).splitlines()

            for step in steps:
                step_str = str(step).strip()
                if not step_str:
                    continue
                p = doc.add_paragraph()
                p.alignment = sec_align
                p.paragraph_format.line_spacing = line_spacing
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.left_indent = Inches(0.25)

                match = re.match(r"^(\d+\.)\s*(.*)$", step_str)
                if match:
                    num_run = p.add_run(match.group(1) + " ")
                    num_run.font.name = font_family
                    num_run.font.size = Pt(content_size)
                    num_run.bold = True

                    text_run = p.add_run(match.group(2))
                    text_run.font.name = font_family
                    text_run.font.size = Pt(content_size)
                else:
                    text_run = p.add_run(step_str)
                    text_run.font.name = font_family
                    text_run.font.size = Pt(content_size)

        elif key == "PROGRAM":
            add_heading(title_text)
            code_text = str(content_val).strip()
            # Render cleanly directly on page without any table or gray box
            p = doc.add_paragraph()
            p.alignment = sec_align
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(content_size - 1)

        elif key == "OUTPUT":
            add_heading(title_text)
            out_text = str(content_val).strip()
            # Render cleanly directly on page without any dark box
            p = doc.add_paragraph()
            p.alignment = sec_align
            p.paragraph_format.line_spacing = line_spacing
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(out_text)
            run.font.name = font_family
            run.font.size = Pt(content_size)

        elif key == "RESULT":
            add_heading(title_text)
            p = doc.add_paragraph()
            p.alignment = sec_align
            p.paragraph_format.line_spacing = line_spacing
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(str(content_val).strip())
            run.font.name = font_family
            run.font.size = Pt(content_size)

        elif key in ("FIGURES", "IMAGES"):
            if images:
                for idx, img_info in enumerate(images, start=1):
                    img_data = img_info.get("data")
                    img_path = img_info.get("path")
                    caption = img_info.get("caption", f"Figure {idx}: Program Execution Console Output")
                    width_in = float(img_info.get("width_in", 4.5))

                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_before = Pt(8)
                        p_img.paragraph_format.space_after = Pt(2)

                        if img_data and img_data.startswith("data:image"):
                            _, b64_str = img_data.split(",", 1)
                            img_bytes = base64.b64decode(b64_str)
                            image_stream = io.BytesIO(img_bytes)
                            p_img.add_run().add_picture(image_stream, width=Inches(width_in))
                        elif img_path and os.path.exists(img_path):
                            p_img.add_run().add_picture(img_path, width=Inches(width_in))

                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_after = Pt(10)
                        run_cap = p_cap.add_run(caption)
                        run_cap.font.name = font_family
                        run_cap.font.size = Pt(content_size - 1.5)
                        run_cap.italic = True
                    except Exception as img_err:
                        print("Failed to add image to document:", img_err)

    # If images were not placed in ordered_items, append them at the end
    has_image_key = any(k in ("FIGURES", "IMAGES") for k, _, _, _ in ordered_items)
    if not has_image_key and images:
        for idx, img_info in enumerate(images, start=1):
            img_data = img_info.get("data")
            img_path = img_info.get("path")
            caption = img_info.get("caption", f"Figure {idx}: Program Execution Console Output")
            width_in = float(img_info.get("width_in", 4.5))

            try:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)

                if img_data and img_data.startswith("data:image"):
                    _, b64_str = img_data.split(",", 1)
                    img_bytes = base64.b64decode(b64_str)
                    image_stream = io.BytesIO(img_bytes)
                    p_img.add_run().add_picture(image_stream, width=Inches(width_in))
                elif img_path and os.path.exists(img_path):
                    p_img.add_run().add_picture(img_path, width=Inches(width_in))

                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(10)
                run_cap = p_cap.add_run(caption)
                run_cap.font.name = font_family
                run_cap.font.size = Pt(content_size - 1.5)
                run_cap.italic = True
            except Exception as img_err:
                print("Failed to add image to document:", img_err)

    # 6. Academic Footer & Evaluator Signature Block
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(36)
    p_sig.paragraph_format.space_after = Pt(4)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    r_line = p_sig.add_run("....................................................\n")
    r_line.font.name = font_family
    r_line.bold = True

    r_sig = p_sig.add_run("SIGNATURE OF THE EVALUATOR")
    r_sig.font.name = font_family
    r_sig.font.size = Pt(content_size)
    r_sig.bold = True

    # Ensure parent directory exists and save document
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    doc.save(output_path)
    return output_path
